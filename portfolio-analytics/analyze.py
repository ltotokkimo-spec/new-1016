"""台股 10 檔權值股投資組合分析：配置最佳化、風險績效指標、相關性與集中度。

資料：weekly_prices.csv（證交所還原權息週收盤，2025-11-07 ~ 2026-10-16，49 週）
輸出：portfolio_analytics_2026-10-16.xlsx、results.json
"""
import csv
import json

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import minimize

NAMES = {
    "2330": "台積電", "2317": "鴻海", "2454": "聯發科", "2382": "廣達", "2308": "台達電",
    "2881": "富邦金", "2891": "中信金", "2412": "中華電", "2303": "聯電", "3711": "日月光投控",
}
BENCH = "0050"
PERIODS = 52
RF = 0.017          # 年化無風險利率（台灣一年期定存約 1.7%）
MAX_W = 0.30        # 最佳化時單一個股權重上限
VAR_LEVEL = 0.95


def load(path="weekly_prices.csv"):
    with open(path, newline="") as f:
        rows = list(csv.reader(f))
    header, data = rows[0], rows[1:]
    dates = [r[0] for r in data]
    prices = np.array([[float(x) for x in r[1:]] for r in data])
    return dates, header[1:], prices


# ---------- 績效與風險指標 ----------
def ann_return(r):
    return float(np.prod(1 + r) ** (PERIODS / len(r)) - 1)


def ann_vol(r):
    return float(np.std(r, ddof=1) * np.sqrt(PERIODS))


def sharpe(r):
    return (ann_return(r) - RF) / ann_vol(r)


def sortino(r):
    rf_p = RF / PERIODS
    downside = np.minimum(r - rf_p, 0)
    dd = np.sqrt(np.mean(downside ** 2)) * np.sqrt(PERIODS)
    return (ann_return(r) - RF) / dd if dd > 0 else float("nan")


def max_drawdown(r):
    wealth = np.cumprod(1 + r)
    peak = np.maximum.accumulate(np.concatenate([[1.0], wealth]))[1:]
    return float(np.min(wealth / peak - 1))


def calmar(r):
    mdd = max_drawdown(r)
    return ann_return(r) / abs(mdd) if mdd < 0 else float("nan")


def beta(r, b):
    return float(np.cov(r, b, ddof=1)[0, 1] / np.var(b, ddof=1))


def hist_var_cvar(r, level=VAR_LEVEL):
    q = np.quantile(r, 1 - level)
    return float(-q), float(-r[r <= q].mean())


def metrics(r, b):
    var, cvar = hist_var_cvar(r)
    return {
        "年化報酬": ann_return(r), "年化波動": ann_vol(r), "Sharpe": sharpe(r),
        "Sortino": sortino(r), "最大回撤": max_drawdown(r), "Calmar": calmar(r),
        "Beta(0050)": beta(r, b), "週VaR95%": var, "週CVaR95%": cvar,
    }


# ---------- 配置最佳化（SLSQP，long-only，單檔上限 MAX_W） ----------
def _solve(obj, n):
    w0 = np.full(n, 1 / n)
    res = minimize(obj, w0, method="SLSQP", bounds=[(0, MAX_W)] * n,
                   constraints=[{"type": "eq", "fun": lambda w: w.sum() - 1}],
                   options={"ftol": 1e-12, "maxiter": 1000})
    w = np.clip(res.x, 0, None)
    return w / w.sum()


def min_variance(cov):
    return _solve(lambda w: w @ cov @ w, len(cov))


def max_sharpe(mu, cov):
    return _solve(lambda w: -(w @ mu - RF) / np.sqrt(w @ cov @ w), len(cov))


def risk_parity(cov):
    n = len(cov)

    def obj(w):
        rc = w * (cov @ w)
        return np.sum((rc - rc.mean()) ** 2) * 1e4

    res = minimize(obj, np.full(n, 1 / n), method="SLSQP", bounds=[(1e-6, 1)] * n,
                   constraints=[{"type": "eq", "fun": lambda w: w.sum() - 1}],
                   options={"ftol": 1e-14, "maxiter": 2000})
    return res.x / res.x.sum()


def risk_contrib(w, cov):
    rc = w * (cov @ w)
    return rc / rc.sum()


def diversification_ratio(w, cov):
    return float(w @ np.sqrt(np.diag(cov)) / np.sqrt(w @ cov @ w))


def hhi(w):
    return float(np.sum(w ** 2))


# ---------- 相關性分群（單一連結階層式，距離 = sqrt(2(1-ρ))） ----------
def correlation_clusters(corr, codes, threshold=0.5):
    """以相關係數 >= threshold 為同群（single-linkage）。"""
    n = len(codes)
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for i in range(n):
        for j in range(i + 1, n):
            if corr[i, j] >= threshold:
                parent[find(i)] = find(j)
    groups = {}
    for i in range(n):
        groups.setdefault(find(i), []).append(codes[i])
    return sorted(groups.values(), key=len, reverse=True)


def main():
    dates, cols, prices = load()
    rets = prices[1:] / prices[:-1] - 1
    bi = cols.index(BENCH)
    codes = [c for c in cols if c != BENCH]
    idx = [cols.index(c) for c in codes]
    R, b = rets[:, idx], rets[:, bi]

    mu = np.array([ann_return(R[:, i]) for i in range(R.shape[1])])
    cov = np.cov(R, rowvar=False, ddof=1) * PERIODS
    corr = np.corrcoef(R, rowvar=False)

    n = len(codes)
    portfolios = {
        "等權重": np.full(n, 1 / n),
        "最小變異": min_variance(cov),
        "最大Sharpe": max_sharpe(mu, cov),
        "風險平價": risk_parity(cov),
    }

    out = {
        "period": [dates[0], dates[-1]], "weeks": len(R), "rf": RF, "max_weight": MAX_W,
        "stocks": {c: metrics(R[:, i], b) for i, c in enumerate(codes)},
        "benchmark": metrics(b, b),
        "portfolios": {}, "corr": corr.tolist(), "codes": codes,
        "clusters": correlation_clusters(corr, codes),
    }
    for name, w in portfolios.items():
        pr = R @ w  # 每週再平衡
        m = metrics(pr, b)
        m.update({"HHI": hhi(w), "有效檔數": 1 / hhi(w),
                  "分散比率": diversification_ratio(w, cov)})
        out["portfolios"][name] = {
            "weights": dict(zip(codes, w.tolist())),
            "risk_contrib": dict(zip(codes, risk_contrib(w, cov).tolist())),
            "metrics": m,
        }

    with open("results.json", "w") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    write_xlsx(out, "portfolio_analytics_2026-10-16.xlsx")
    return out


HDR_FILL = PatternFill("solid", fgColor="305496")
PCT_KEYS = {"年化報酬", "年化波動", "最大回撤", "週VaR95%", "週CVaR95%"}


def _header(ws, row):
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def _fmt_row(ws, keys, start_col=2):
    r = ws.max_row
    for j, k in enumerate(keys):
        ws.cell(r, start_col + j).number_format = "0.0%" if k in PCT_KEYS else "0.00"


def write_xlsx(out, path):
    codes = out["codes"]
    label = [f"{c} {NAMES[c]}" for c in codes]
    wb = Workbook()

    ws = wb.active
    ws.title = "配置比較"
    names = list(out["portfolios"])
    _header(ws, ["股票"] + [f"{p} 權重" for p in names] + [f"{p} 風險貢獻" for p in names])
    for i, c in enumerate(codes):
        ws.append([label[i]] + [out["portfolios"][p]["weights"][c] for p in names]
                  + [out["portfolios"][p]["risk_contrib"][c] for p in names])
        for j in range(2, 2 + 2 * len(names)):
            ws.cell(ws.max_row, j).number_format = "0.0%"
    ws.append([])
    mkeys = list(out["portfolios"][names[0]]["metrics"])
    _header(ws, ["指標"] + names + ["0050 基準"])
    for k in mkeys:
        bench = out["benchmark"].get(k, "")
        ws.append([k] + [out["portfolios"][p]["metrics"][k] for p in names] + [bench])
        for j in range(2, 3 + len(names)):
            ws.cell(ws.max_row, j).number_format = "0.0%" if k in PCT_KEYS else "0.00"
    ws.column_dimensions["A"].width = 16
    for j in range(2, 2 + 2 * len(names)):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 13

    ws = wb.create_sheet("個股指標")
    keys = list(out["benchmark"])
    _header(ws, ["股票"] + keys)
    for i, c in enumerate(codes):
        ws.append([label[i]] + [out["stocks"][c][k] for k in keys])
        _fmt_row(ws, keys)
    ws.append(["0050 元大台灣50"] + [out["benchmark"][k] for k in keys])
    _fmt_row(ws, keys)
    ws.column_dimensions["A"].width = 16
    for j in range(2, 2 + len(keys)):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 11

    ws = wb.create_sheet("相關係數")
    _header(ws, [""] + label)
    for i in range(len(codes)):
        ws.append([label[i]] + out["corr"][i])
        for j in range(2, 2 + len(codes)):
            ws.cell(ws.max_row, j).number_format = "0.00"
    rng = f"B2:{ws.cell(1 + len(codes), 1 + len(codes)).coordinate}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=-0.2, start_color="63BE7B",
        mid_type="num", mid_value=0.3, mid_color="FFFFFF",
        end_type="num", end_value=1, end_color="F8696B"))
    ws.column_dimensions["A"].width = 16
    ws.append([])
    ws.append(["相關分群（ρ ≥ 0.5 視為同群）"])
    for g in out["clusters"]:
        ws.append(["、".join(f"{c} {NAMES[c]}" for c in g)])

    ws = wb.create_sheet("說明")
    for line in [
        ["期間", f"{out['period'][0]} ~ {out['period'][1]}，{out['weeks']} 個週報酬"],
        ["資料", "證交所還原權息日K，取每週最後交易日收盤（含息報酬）"],
        ["年化", "報酬為幾何年化，波動以週標準差 × √52"],
        ["無風險利率", f"{out['rf']:.1%}"],
        ["最佳化", f"SLSQP、只做多、單檔上限 {out['max_weight']:.0%}、權重和 = 100%；每週再平衡"],
        ["VaR/CVaR", "歷史模擬法，95% 信賴水準，單位為週報酬"],
        ["限制", "樣本只有 49 週，且涵蓋 AI 大多頭；歷史平均報酬對未來的預測力很弱，"
                 "最大 Sharpe 權重對輸入非常敏感，請勿直接當作買賣依據"],
        ["免責聲明", "僅供研究參考，不構成投資建議"],
    ]:
        ws.append(line)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 100
    wb.save(path)


if __name__ == "__main__":
    o = main()
    for p, d in o["portfolios"].items():
        w = ", ".join(f"{NAMES[c]} {v:.0%}" for c, v in d["weights"].items() if v > 0.005)
        m = d["metrics"]
        print(f"{p}: 報酬 {m['年化報酬']:.1%} 波動 {m['年化波動']:.1%} Sharpe {m['Sharpe']:.2f} "
              f"MDD {m['最大回撤']:.1%} Beta {m['Beta(0050)']:.2f} 有效檔數 {m['有效檔數']:.1f} 分散比 {m['分散比率']:.2f}")
        print("   ", w)
    bm = o["benchmark"]
    print(f"0050: 報酬 {bm['年化報酬']:.1%} 波動 {bm['年化波動']:.1%} Sharpe {bm['Sharpe']:.2f} MDD {bm['最大回撤']:.1%}")
    print()
    for c, m in o["stocks"].items():
        print(f"{c} {NAMES[c]:5s} 報酬 {m['年化報酬']:7.1%} 波動 {m['年化波動']:6.1%} Sharpe {m['Sharpe']:5.2f} "
              f"Sortino {m['Sortino']:5.2f} MDD {m['最大回撤']:6.1%} Beta {m['Beta(0050)']:5.2f} "
              f"VaR {m['週VaR95%']:5.1%} CVaR {m['週CVaR95%']:5.1%}")
    print()
    print("風險貢獻（等權重）:", ", ".join(f"{NAMES[c]} {v:.0%}" for c, v in o["portfolios"]["等權重"]["risk_contrib"].items()))
    print("分群:", o["clusters"])
    corr = np.array(o["corr"])
    iu = np.triu_indices(len(o["codes"]), 1)
    print(f"平均兩兩相關 {corr[iu].mean():.2f}")
    pairs = sorted(zip(corr[iu], iu[0], iu[1]))
    for v, i, j in pairs[-3:] + pairs[:3]:
        print(f"  {NAMES[o['codes'][i]]}-{NAMES[o['codes'][j]]}: {v:.2f}")
